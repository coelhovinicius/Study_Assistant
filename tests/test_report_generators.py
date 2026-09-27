"""Testes dos geradores de relatório final (garantem que cada formato
produz bytes não vazios e não estoura exceção com um conteúdo típico)."""

from __future__ import annotations

import io
from datetime import datetime, timezone

from study_assistant.domain.entities import (
    AnalysisResult,
    ApostilaInsights,
    ExtractedSection,
    ExtractionMethod,
    Material,
    MaterialType,
    ProviderAttempt,
    SourceFormat,
    StudySession,
)
from study_assistant.infrastructure.report_generators.csv_generator import CsvReportGenerator
from study_assistant.infrastructure.report_generators.docx_generator import DocxReportGenerator
from study_assistant.infrastructure.report_generators.pdf_generator import PdfReportGenerator
from study_assistant.infrastructure.report_generators.txt_generator import TxtReportGenerator


def _sample_session() -> StudySession:
    material = Material(
        filename="apostila.pdf",
        material_type=MaterialType.APOSTILA,
        source_format=SourceFormat.PDF,
        raw_text="conteúdo de exemplo da apostila",
    )
    insights = ApostilaInsights(
        referencias_bibliograficas=ExtractedSection("MENDES, Gilmar...", ExtractionMethod.HEURISTICA),
        dicas_leitura=ExtractedSection("Leia também...", ExtractionMethod.HEURISTICA),
        desafio_pratico=ExtractedSection("Resolva o caso X", ExtractionMethod.HEURISTICA),
    )
    analysis = AnalysisResult(
        sections={
            "analise_apostila": "Análise detalhada da apostila com acentuação: ção, ã, é.",
            "analise_e_resolucao_desafio": "Resolução completa do desafio proposto.",
        },
        generated_by_provider="OpenAI",
        generated_by_model="gpt-5.5",
        provider_attempts=(
            ProviderAttempt("OpenAI", "gpt-5.5", success=True, duration_ms=123.4),
        ),
    )
    return StudySession(
        title="Sessão de teste — Direito Constitucional",
        materials=[material],
        apostila_insights=insights,
        analysis_result=analysis,
    )


def test_txt_generator_produz_bytes_nao_vazios() -> None:
    content = TxtReportGenerator().generate(_sample_session())
    assert isinstance(content, bytes)
    assert "Análise detalhada".encode("utf-8") in content
    assert "Resolução completa".encode("utf-8") in content


def test_csv_generator_produz_bytes_nao_vazios() -> None:
    content = CsvReportGenerator().generate(_sample_session())
    assert isinstance(content, bytes)
    assert len(content) > 0


def test_csv_generator_neutraliza_conteudo_que_pareceria_formula() -> None:
    """Se a IA (ou o material enviado) gerar um texto que comece com =, +,
    -, @ ou tab, o Excel/Sheets podem interpretar a célula como fórmula ao
    abrir o CSV — o gerador precisa prefixar com aspas simples pra evitar
    isso (mitigação padrão de "CSV injection")."""
    session = _sample_session()
    session.analysis_result.sections["analise_apostila"] = "=cmd|'/c calc'!A1"

    content = CsvReportGenerator().generate(session).decode("utf-8-sig")

    assert "'=cmd|'/c calc'!A1" in content
    # a fórmula "crua" (sem o prefixo de aspas) não pode aparecer sozinha
    # logo depois de uma aspa de campo CSV, senão o Excel ainda a executaria
    assert ',"=cmd' not in content


def test_docx_generator_produz_bytes_nao_vazios() -> None:
    content = DocxReportGenerator().generate(_sample_session())
    assert isinstance(content, bytes)
    assert content[:2] == b"PK"  # docx é um zip


def test_pdf_generator_produz_bytes_nao_vazios() -> None:
    content = PdfReportGenerator().generate(_sample_session())
    assert isinstance(content, bytes)
    assert content.startswith(b"%PDF")


def test_txt_generator_nao_menciona_mais_o_provedor_de_ia() -> None:
    """O usuário pediu pra tirar os avisos de 'gerado por IA' do relatório
    — a linha "Provedor de IA: ..." não pode mais aparecer."""
    content = TxtReportGenerator().generate(_sample_session()).decode("utf-8")
    assert "Provedor de IA" not in content
    assert "OpenAI" not in content  # o nome do provedor também não deveria mais vazar


def test_docx_generator_nao_menciona_mais_o_provedor_e_justifica_o_corpo() -> None:
    import io

    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH

    content = DocxReportGenerator().generate(_sample_session())
    document = Document(io.BytesIO(content))

    full_text = "\n".join(p.text for p in document.paragraphs)
    assert "Provedor de IA" not in full_text
    assert "OpenAI" not in full_text

    body_paragraphs = [p for p in document.paragraphs if "Análise detalhada" in p.text]
    assert body_paragraphs, "parágrafo de análise não encontrado no docx gerado"
    assert body_paragraphs[0].alignment == WD_ALIGN_PARAGRAPH.JUSTIFY


def test_pdf_generator_nao_menciona_mais_o_provedor_e_usa_horario_de_brasilia() -> None:
    from pypdf import PdfReader

    session = _sample_session()
    session.created_at = datetime(2026, 8, 31, 14, 30, tzinfo=timezone.utc)
    content = PdfReportGenerator().generate(session)

    reader = PdfReader(io.BytesIO(content))
    text = "\n".join(page.extract_text() or "" for page in reader.pages)

    assert "Provedor de IA" not in text
    assert "OpenAI" not in text
    # 14:30 UTC == 11:30 em Brasília (UTC-3) — confirma a conversão de fuso
    # sem exigir o rótulo "(horário de Brasília)", que foi removido por ser
    # redundante pro usuário (ele já sabe que é o fuso dele).
    assert "31/08/2026 11:30" in text
    assert "horário de Brasília" not in text
    assert "Análise detalhada" in text


def _session_with_parts() -> StudySession:
    session = _sample_session()
    session.analysis_result = AnalysisResult(
        sections={
            "analise_apostila": (
                "### Parte 1 de 2 — Fundamentos\nAnálise da primeira parte.\n\n"
                "### Parte 2 de 2 — Organização do Estado\nAnálise da segunda parte."
            )
        },
        generated_by_provider="n8n",
        generated_by_model="cascata",
        provider_attempts=(),
    )
    return session


def test_docx_mostra_as_partes_da_analise_como_subtitulo() -> None:
    from docx import Document

    document = Document(io.BytesIO(DocxReportGenerator().generate(_session_with_parts())))

    headings = [p.text for p in document.paragraphs if p.style.name == "Heading 2"]
    assert "Parte 1 de 2 — Fundamentos" in headings
    assert "Parte 2 de 2 — Organização do Estado" in headings
    assert not any("###" in p.text for p in document.paragraphs)


def test_pdf_mostra_as_partes_da_analise_sem_o_marcador() -> None:
    from pypdf import PdfReader

    reader = PdfReader(io.BytesIO(PdfReportGenerator().generate(_session_with_parts())))
    text = "\n".join(page.extract_text() for page in reader.pages)

    assert "Parte 2 de 2 — Organização do Estado" in text
    assert "###" not in text


def test_pdf_nao_desenha_quadrado_preto_no_lugar_de_hifen() -> None:
    """Regressão: a IA escreve "diferenciando‑se" com o hífen que não quebra
    linha (U+2011), que a Helvetica não tem — o PDF saía com um quadrado
    preto no lugar de cada hífen."""
    from pypdf import PdfReader

    session = _sample_session()
    session.analysis_result = AnalysisResult(
        sections={
            "analise_apostila": (
                "Diferenciando\u2011se do hardware, a arquitetura cliente\u2010servidor "
                "usa deploy blue\u2011green\u200b e marcadores \u25aa \U0001f680 fim."
            )
        },
        generated_by_provider="n8n",
        generated_by_model="cascata",
        provider_attempts=(),
    )

    reader = PdfReader(io.BytesIO(PdfReportGenerator().generate(session)))
    text = "\n".join(page.extract_text() for page in reader.pages)

    assert "\u25a0" not in text  # ■
    assert "Diferenciando-se" in text
    assert "cliente-servidor" in text
    assert "blue-green" in text


def _session_with_bold() -> StudySession:
    session = _sample_session()
    session.analysis_result = AnalysisResult(
        sections={"analise_apostila": "Um **Método** formaliza as atividades de um **Processo**."},
        generated_by_provider="n8n",
        generated_by_model="cascata",
        provider_attempts=(),
    )
    return session


def test_pdf_aplica_o_negrito_da_ia_sem_mostrar_asteriscos() -> None:
    from pypdf import PdfReader

    reader = PdfReader(io.BytesIO(PdfReportGenerator().generate(_session_with_bold())))
    page = reader.pages[0]
    text = page.extract_text()
    fonts = {str(f.get_object()["/BaseFont"]) for f in page["/Resources"]["/Font"].values()}

    assert "**" not in text
    assert "Um Método formaliza as atividades de um Processo." in " ".join(text.split())
    assert "/Helvetica-Bold" in fonts


def test_docx_aplica_o_negrito_da_ia_sem_mostrar_asteriscos() -> None:
    from docx import Document

    document = Document(io.BytesIO(DocxReportGenerator().generate(_session_with_bold())))
    paragraph = next(p for p in document.paragraphs if "formaliza" in p.text)

    assert paragraph.text == "Um Método formaliza as atividades de um Processo."
    assert [run.text for run in paragraph.runs if run.bold] == ["Método", "Processo"]
