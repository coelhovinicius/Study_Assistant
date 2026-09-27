"""Gerador de relatório em Word (.docx), usando python-docx."""

from __future__ import annotations

import io

from study_assistant.domain.entities import SECTION_SUBHEADING_PREFIX, StudySession
from study_assistant.domain.exceptions import ReportGenerationError
from study_assistant.domain.ports import ReportGenerator
from study_assistant.infrastructure.report_generators.report_content import build_report_outline


class DocxReportGenerator(ReportGenerator):
    @property
    def format_name(self) -> str:
        return "docx"

    @property
    def mime_type(self) -> str:
        return "application/vnd.openxmlformats-officedocument.wordprocessingml.document"

    def generate(self, session: StudySession) -> bytes:
        try:
            import docx
        except ImportError as exc:  # pragma: no cover
            raise ReportGenerationError(
                "Dependência 'python-docx' não instalada. Rode: pip install python-docx"
            ) from exc

        from docx.enum.text import WD_ALIGN_PARAGRAPH

        outline = build_report_outline(session)
        document = docx.Document()

        document.add_heading(outline.title, level=0)
        meta = document.add_paragraph()
        meta.add_run(f"Gerado em: {outline.generated_at_label}").italic = True

        if outline.materials:
            document.add_heading("Materiais Analisados", level=1)
            for material in outline.materials:
                document.add_paragraph(
                    f"{material.filename} — {material.type_label}", style="List Bullet"
                )

        if outline.insights:
            document.add_heading("Extrações da Apostila", level=1)
            for insight in outline.insights:
                document.add_heading(f"{insight.label}", level=2)
                document.add_paragraph(f"({insight.method_label})").italic = True
                body = document.add_paragraph(insight.content)
                body.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

        for title, content in outline.sections:
            document.add_heading(title, level=1)
            for paragraph_text in content.split("\n"):
                if not paragraph_text.strip():
                    continue
                if paragraph_text.startswith(SECTION_SUBHEADING_PREFIX):
                    # "### Parte 2 de 4 — ..." (análise em lotes) sai como subtítulo.
                    document.add_heading(paragraph_text[len(SECTION_SUBHEADING_PREFIX):].strip(), level=2)
                    continue
                body = document.add_paragraph(paragraph_text)
                body.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

        buffer = io.BytesIO()
        document.save(buffer)
        return buffer.getvalue()
