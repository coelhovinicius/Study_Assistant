"""Extração de texto de arquivos DOCX (Word)."""

from __future__ import annotations

import io

from study_assistant.domain.exceptions import EmptyDocumentError, TextExtractionError
from study_assistant.domain.ports import TextExtractor


class DocxTextExtractor(TextExtractor):
    """Extrai texto de documentos .docx, incluindo texto de tabelas."""

    def supports(self, filename: str) -> bool:
        return filename.lower().endswith(".docx")

    def extract(self, file_bytes: bytes, filename: str) -> str:
        try:
            import docx  # python-docx
        except ImportError as exc:  # pragma: no cover
            raise TextExtractionError(
                "Dependência 'python-docx' não instalada. Rode: pip install python-docx"
            ) from exc

        try:
            document = docx.Document(io.BytesIO(file_bytes))
            parts = [p.text for p in document.paragraphs if p.text.strip()]
            for table in document.tables:
                for row in table.rows:
                    row_text = " | ".join(cell.text.strip() for cell in row.cells)
                    if row_text.strip(" |"):
                        parts.append(row_text)
            text = "\n".join(parts).strip()
        except Exception as exc:  # noqa: BLE001
            raise TextExtractionError(
                f"Falha ao extrair texto do DOCX '{filename}': {exc}"
            ) from exc

        if not text:
            raise EmptyDocumentError(f"O documento '{filename}' não contém texto.")
        return text
